from __future__ import annotations

import re
import unicodedata
from pathlib import Path

# ---------------------------------------------------------------------------
# CJK-aware tokenization
#
# SQLite's unicode61 tokenizer does NOT split Chinese: "设备采购方案" becomes one
# token, so searching "采购" finds nothing. We therefore store a pre-tokenized
# copy of every field (CJK -> one token per character, Latin -> words) and turn
# user queries into FTS5 phrases such as "采 购" (adjacent characters).
# ---------------------------------------------------------------------------

_CJK_RANGES = (
    (0x3400, 0x4DBF),
    (0x4E00, 0x9FFF),
    (0xF900, 0xFAFF),
    (0x3040, 0x30FF),
    (0xAC00, 0xD7AF),
)


def is_cjk(ch: str) -> bool:
    cp = ord(ch)
    return any(lo <= cp <= hi for lo, hi in _CJK_RANGES)


def normalize(text: str) -> str:
    """NFKC (full-width -> half-width), lower-case, collapse whitespace."""
    if not text:
        return ""
    text = unicodedata.normalize("NFKC", text)
    text = text.lower()
    return re.sub(r"\s+", " ", text).strip()


def tokenize_cjk(text: str) -> str:
    """Space-separated tokens: one per CJK char, one per Latin/digit word."""
    out: list[str] = []
    buf: list[str] = []

    def flush() -> None:
        if buf:
            out.append("".join(buf))
            buf.clear()

    for ch in text:
        if is_cjk(ch):
            flush()
            out.append(ch)
        elif ch.isalnum():
            buf.append(ch)
        else:
            flush()
    flush()
    return " ".join(out)


def _runs(text: str):
    """Yield (substring, is_cjk) for maximal runs of CJK / non-CJK."""
    if not text:
        return
    buf: list[str] = []
    cur: bool | None = None
    for ch in text:
        flag = is_cjk(ch)
        if cur is None:
            cur = flag
        if flag != cur:
            yield "".join(buf), cur
            buf = []
            cur = flag
        buf.append(ch)
    if buf:
        yield "".join(buf), cur


def build_fts_query(query: str, mode: str = "and") -> str:
    """Natural query -> FTS5 expression.

    mode="and": whole CJK run as one phrase ("报 价 单") + Latin prefixes.
    mode="or" : every CJK bigram as a phrase ("报 价" OR "价 单") + Latin prefixes.
                Used as a recall fallback for long natural-language queries,
                where a strict AND almost never matches ("太高" vs "过高").
    """
    q = normalize(query)
    if not q:
        return ""
    terms: list[str] = []
    for chunk, cjk in _runs(q):
        if cjk:
            toks = [c for c in chunk if is_cjk(c)]
            if not toks:
                continue
            if len(toks) == 1:
                terms.append('"' + toks[0] + '"')
            elif mode == "and":
                terms.append('"' + " ".join(toks) + '"')
            else:
                for i in range(len(toks) - 1):
                    terms.append(f'"{toks[i]} {toks[i + 1]}"')
        else:
            for word in re.findall(r"[a-z0-9]+", chunk):
                if len(word) >= 1:
                    terms.append(word + "*")
    seen: set[str] = set()
    uniq = []
    for t in terms:
        if t not in seen:
            seen.add(t)
            uniq.append(t)
    joiner = " AND " if mode == "and" else " OR "
    return joiner.join(uniq)


# ---------------------------------------------------------------------------
# Chunking
# ---------------------------------------------------------------------------

_SENT_SPLIT = re.compile(r"(?<=[。！？!?；;])|\n{2,}")


def chunk_text(text: str, size: int = 400, overlap: int = 60, max_chunks: int = 12) -> list[str]:
    text = re.sub(r"[ \t]+", " ", text or "").strip()
    if not text:
        return []
    if len(text) <= size:
        return [text]
    sentences = [s.strip() for s in _SENT_SPLIT.split(text) if s and s.strip()]
    chunks: list[str] = []
    cur = ""
    for s in sentences:
        if len(cur) + len(s) + 1 <= size:
            cur = (cur + " " + s).strip()
            continue
        if cur:
            chunks.append(cur)
        # a single sentence longer than size: hard split
        while len(s) > size:
            chunks.append(s[:size])
            s = s[size - overlap :]
        cur = s
        if len(chunks) >= max_chunks:
            break
    if cur and len(chunks) < max_chunks:
        chunks.append(cur)
    return chunks[:max_chunks]


# ---------------------------------------------------------------------------
# Text extraction (optional dependencies degrade gracefully)
# ---------------------------------------------------------------------------

_TEXT_EXTS = {
    ".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".jsonl", ".log",
    ".py", ".js", ".ts", ".tsx", ".jsx", ".java", ".kt", ".c", ".h", ".cpp",
    ".hpp", ".cs", ".go", ".rs", ".rb", ".php", ".sh", ".ps1", ".sql", ".yml",
    ".yaml", ".toml", ".ini", ".cfg", ".xml", ".html", ".htm", ".rst",
}


def _read_text_file(path: Path, max_bytes: int = 2_000_000) -> str:
    raw = path.read_bytes()[:max_bytes]
    for enc in ("utf-8", "gb18030", "utf-16", "latin-1"):
        try:
            return raw.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return raw.decode("utf-8", "ignore")


def _extract_pdf(path: Path, max_pages: int, max_chars: int) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        return ""
    try:
        reader = PdfReader(str(path))
        parts: list[str] = []
        for page in reader.pages[:max_pages]:
            parts.append(page.extract_text() or "")
            if sum(len(p) for p in parts) >= max_chars:
                break
        return "\n".join(parts)[:max_chars]
    except Exception:
        return ""


def _extract_docx(path: Path, max_chars: int) -> str:
    try:
        import docx
    except ImportError:
        return ""
    try:
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs if p.text]
        for table in d.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        parts.append(cell.text)
        return "\n".join(parts)[:max_chars]
    except Exception:
        return ""


def _extract_xlsx(path: Path, max_chars: int) -> str:
    try:
        import openpyxl
    except ImportError:
        return ""
    try:
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        parts: list[str] = []
        for ws in wb.worksheets[:4]:
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    parts.append(" ".join(cells))
                if sum(len(p) for p in parts) >= max_chars:
                    break
        wb.close()
        return "\n".join(parts)[:max_chars]
    except Exception:
        return ""


def extract_text(
    path: str | Path,
    *,
    max_chars: int = 4000,
    max_pages: int = 8,
    max_bytes: int = 2_000_000,
) -> str:
    p = Path(path)
    ext = p.suffix.lower()
    if ext in _TEXT_EXTS:
        return _read_text_file(p, max_bytes)[:max_chars]
    if ext == ".pdf":
        return _extract_pdf(p, max_pages, max_chars)
    if ext == ".docx":
        return _extract_docx(p, max_chars)
    if ext in (".xlsx", ".xlsm"):
        return _extract_xlsx(p, max_chars)
    return ""


# ---------------------------------------------------------------------------
# Aliases and labels
# ---------------------------------------------------------------------------


def pinyin_aliases(text: str, max_chars: int = 24) -> str:
    """Full pinyin + initials, e.g. 报价单 -> 'baojiadan bjd'."""
    try:
        from pypinyin import Style, lazy_pinyin
    except ImportError:
        return ""
    core = "".join(ch for ch in text if is_cjk(ch))[:max_chars]
    if not core:
        return ""
    full = "".join(lazy_pinyin(core))
    initials = "".join(lazy_pinyin(core, style=Style.FIRST_LETTER))
    return f"{full} {initials}".strip()


def size_label(n: int | None) -> str:
    if n is None:
        return ""
    units = ["B", "KB", "MB", "GB", "TB"]
    v = float(n)
    for u in units:
        if v < 1024 or u == units[-1]:
            return f"{v:.0f}{u}" if u == "B" else f"{v:.1f}{u}"
        v /= 1024
    return f"{n}B"


def date_label(ts: float | None) -> str:
    if not ts:
        return ""
    import datetime as _dt

    d = _dt.datetime.fromtimestamp(ts)
    return f"{d.year}年{d.month}月{d.day}日"
